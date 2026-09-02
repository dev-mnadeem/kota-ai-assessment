"""Build and validate immutable dataset snapshots before publishing them."""
from __future__ import annotations

import csv
from contextlib import closing
import fcntl
import hashlib
import io
import json
import os
import re
from pathlib import Path
import sqlite3
import tempfile
import time
from datetime import datetime, timezone
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
from urllib.error import HTTPError
import zipfile

from scripts.normalize_to_sqlite import ensure_schema, iter_data_records

ROOT = Path(__file__).resolve().parents[1]
COLUMNS = (
    'table_id','year','segment','work_pattern','primary_group_type','primary_group_value',
    'earnings_band','pension_type','measure_type','value_num','value_raw','suppression_flag','source_file',
)
MAX_ARCHIVE_BYTES = 5_000_000
MAX_WORKBOOK_BYTES = 20_000_000


def fingerprint(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def status(db_path: str | Path) -> dict:
    path = Path(db_path).resolve()
    if not path.is_file():
        return {'ready': False, 'reason': 'Dataset has not been built'}
    try:
        with closing(sqlite3.connect(path.as_uri() + '?mode=ro', uri=True)) as connection:
            metadata = dict(connection.execute('SELECT key, value FROM dataset_metadata'))
            rows = connection.execute('SELECT COUNT(*) FROM observations').fetchone()[0]
            tables = [r[0] for r in connection.execute('SELECT DISTINCT table_id FROM observations ORDER BY table_id')]
            years = [r[0] for r in connection.execute('SELECT DISTINCT year FROM observations ORDER BY year')]
        return {'ready': rows > 0, 'rows': rows, 'tables': tables, 'years': years, **metadata}
    except sqlite3.Error:
        return {'ready': False, 'reason': 'Dataset schema or provenance is invalid'}


def archive_bytes(source: dict, cache: Path) -> bytes:
    url = source['url']
    parsed = urlsplit(url)
    if parsed.scheme != 'https' or parsed.hostname != 'www.ons.gov.uk':
        raise ValueError('Only pinned HTTPS ONS downloads are accepted')
    expected = source['sha256']
    if not isinstance(expected, str) or not re.fullmatch(r'[0-9a-f]{64}', expected):
        raise ValueError('A SHA-256 source checksum is required')
    path = cache / (expected + '.zip')
    if path.exists():
        data = path.read_bytes()
    else:
        for attempt in range(4):
            time.sleep(2)
            try:
                with urlopen(Request(url, headers={'User-Agent':'ASHE reproducible importer'}), timeout=45) as response:
                    if urlsplit(response.url).hostname != 'www.ons.gov.uk':
                        raise ValueError('Unexpected download redirect')
                    data = response.read(MAX_ARCHIVE_BYTES + 1)
                break
            except HTTPError as error:
                if error.code not in {429,500,502,503,504} or attempt == 3:
                    raise
                retry_after = error.headers.get('Retry-After', '20')
                delay = int(retry_after) if retry_after.isdigit() else 30
                if delay > 300:
                    raise ValueError('Source asked for a long retry delay; try the refresh later') from error
                time.sleep(delay)

    if len(data) > MAX_ARCHIVE_BYTES or hashlib.sha256(data).hexdigest() != expected:
        raise ValueError('Source size or checksum mismatch; the live dataset was not replaced')
    path.write_bytes(data)
    return data


def source_records(source: dict, cache: Path, temporary: Path):
    if source.get('format') == 'normalized-csv':
        path = (ROOT / source['path']).resolve()
        if not path.is_relative_to(ROOT / 'fixtures') or fingerprint(path) != source['sha256']:
            raise ValueError('Fixture path or checksum mismatch')
        with path.open(newline='') as handle:
            for row in csv.DictReader(handle):
                row['year'] = int(row['year'])
                row['value_num'] = float(row['value_num']) if row['value_num'] else None
                yield {key: value if value != '' else None for key, value in row.items()}
        return
    with zipfile.ZipFile(io.BytesIO(archive_bytes(source, cache))) as archive:
        workbooks = [info for info in archive.infolist()
                     if info.filename.lower().endswith(('.xlsx','.xls'))
                     and 'cv' not in info.filename.lower()]
        if not workbooks:
            raise ValueError('No estimate workbook found in the pinned archive')
        for info in workbooks:
            if info.file_size > MAX_WORKBOOK_BYTES:
                raise ValueError('Workbook exceeds the expansion limit')
            raw = archive.read(info)
            if info.filename.lower().endswith('.xlsx'):
                import openpyxl
                workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                try:
                    rows = list(workbook['All'].iter_rows(values_only=True))
                finally:
                    workbook.close()
            else:
                import xlrd
                workbook = xlrd.open_workbook(file_contents=raw)
                sheet = workbook.sheet_by_name('All')
                rows = [sheet.row_values(index) for index in range(sheet.nrows)]
                workbook.release_resources()
            table_id = source['table_id']
            year = int(source['year'])
            if not re.fullmatch(r'P(?:1[0-2]|[1-9])', table_id):
                raise ValueError('Unexpected table identifier')
            prepared = temporary / f"{year}_{table_id}__All.csv"
            with prepared.open('w',newline='') as handle:
                csv.writer(handle).writerows([[str(value) if value is not None else '' for value in row] for row in rows])
            records = iter_data_records(prepared)
            if not records:
                raise ValueError('Pinned source produced no observations')
            for row in records:
                row['source_file'] = source['url'] + '#' + info.filename + '/All'
                yield row


def build(manifest_path: str | Path, db_path: str | Path) -> dict:
    manifest_path = Path(manifest_path).resolve()
    manifest = json.loads(manifest_path.read_text())
    if not manifest.get('sources') or manifest.get('mode') not in {'fixture','ons'}:
        raise ValueError('A versioned source manifest is required')
    target = Path(db_path).resolve()
    target.parent.mkdir(parents=True,exist_ok=True)
    cache = target.parent / 'cache';cache.mkdir(exist_ok=True)
    with (target.parent / 'refresh.lock').open('a') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        with tempfile.TemporaryDirectory(prefix='.ashe-build-', dir=target.parent) as folder:
            temporary = Path(folder)
            staged = temporary / 'snapshot.db'
            with closing(sqlite3.connect(staged)) as connection, connection:
                ensure_schema(connection)
                placeholders = ','.join('?' for _ in COLUMNS)
                for source in manifest['sources']:
                    rows = list(source_records(source, cache, temporary))
                    connection.executemany(f"INSERT INTO observations ({','.join(COLUMNS)}) VALUES ({placeholders})", [tuple(row[key] for key in COLUMNS) for row in rows])
                    connection.execute('INSERT INTO ingest_log (source_file,status,rows_loaded,error_message) VALUES (?,?,?,NULL)', (source.get('url') or source['path'],'ok',len(rows)))
                connection.execute('CREATE TABLE dataset_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)')
                metadata = {'version':manifest['version'], 'mode':manifest['mode'],
                            'manifest_sha256':fingerprint(manifest_path),
                            'built_at':datetime.now(timezone.utc).isoformat(),
                            'notice':manifest.get('notice','Pinned ONS sources; inspect the manifest for editions and coverage.')}
                connection.executemany('INSERT INTO dataset_metadata VALUES (?,?)',metadata.items())
                if connection.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
                    raise ValueError('SQLite integrity check failed')
            result = status(staged)
            if not result['ready']:
                raise ValueError('Dataset validation failed')
            os.replace(staged,target)
            return result
